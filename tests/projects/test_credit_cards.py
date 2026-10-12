"""Credit-card feature (HDFC BizBlack x2): reader, model, policies, template contract.

Policies locked here (Niloy, 2026-10-12):
  * funding order cards > OD > treasury
  * clear-by = use date + (55 - 7) = 48 days, never later than HDFC's bill due date
  * card dues inside the priority window (or overdue) surface; paid (cleared_on) spends never do
  * only last-4 digits are ever stored — a full card number pasted into the sheet is rejected
"""
import json
import sys
import unittest
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts" / "projects"))

import cards as C  # noqa: E402
import read_cashflow_xlsx as R  # noqa: E402

A = dict(C.DEFAULTS)
CARD = {"stmtDay": 15, "dueGapDays": 25}


def raw(cards=None, txns=None, config=None):
    return {"present": True, "config": config or {}, "meta": {"warnings": []},
            "cards": cards if cards is not None else [
                {"card_last4": "9986", "card_holder": "Monica Debnath", "product": "HDFC BizBlack Metal",
                 "credit_limit": 400000, "outstanding": 138000, "statement_day": 15, "due_gap_days": 25,
                 "as_of_date": "2026-10-12", "updated_by": "Sonal", "welcome_spend_target": 150000},
                {"card_last4": "3079", "card_holder": "Niloy Debnath", "product": "HDFC BizBlack Metal",
                 "credit_limit": 400000, "outstanding": 272000, "statement_day": 21, "due_gap_days": 25,
                 "as_of_date": "2026-10-12", "updated_by": "Sonal", "welcome_spend_target": 150000}],
            "txns": txns if txns is not None else []}


class TestClearBy(unittest.TestCase):
    def test_rule_is_use_plus_48_days(self):
        # spend just after statement day -> max float; 48d rule beats bill due
        d, basis = C.clear_by(CARD, date(2026, 10, 16), A)
        self.assertEqual(d, date(2026, 12, 3))
        self.assertEqual(basis, "use + 48d")

    def test_capped_at_bill_due_when_earlier(self):
        # spend 7 Oct, statement 15 Oct, due 9 Nov (33d) < 48d rule
        d, basis = C.clear_by(CARD, date(2026, 10, 7), A)
        self.assertEqual(d, date(2026, 11, 9))
        self.assertEqual(basis, "HDFC bill due")

    def test_bill_due_year_wrap_and_short_month(self):
        self.assertEqual(C.bill_due(15, 25, date(2026, 12, 20)), date(2027, 2, 9))
        # statement day 31 clamps to month end (Feb)
        self.assertEqual(C.bill_due(31, 20, date(2027, 2, 10)), date(2027, 3, 20))


class TestComputeCards(unittest.TestCase):
    TODAY = date(2026, 10, 12)

    def test_none_when_tabs_absent(self):
        self.assertIsNone(C.compute_cards({"present": False, "cards": []}, self.TODAY))
        self.assertIsNone(C.compute_cards(None, self.TODAY))

    def test_totals_and_utilisation(self):
        b = C.compute_cards(raw(), self.TODAY)
        self.assertEqual(b["totals"]["limit"], 800000)
        self.assertEqual(b["totals"]["outstanding"], 410000)
        self.assertEqual(b["totals"]["available"], 390000)
        self.assertEqual(b["totals"]["utilPct"], 51)

    def test_demo_flag_from_updated_by_or_notes(self):
        r = raw()
        r["cards"][0]["updated_by"] = "DEMO"
        self.assertTrue(C.compute_cards(r, self.TODAY)["isDemo"])
        self.assertFalse(C.compute_cards(raw(), self.TODAY)["isDemo"])
        r2 = raw(txns=[{"date": "2026-10-01", "card_last4": "9986", "description": "x", "category": "GST",
                        "amount": 1000, "notes": "DEMO DATA"}])
        self.assertTrue(C.compute_cards(r2, self.TODAY)["isDemo"])

    def test_due_item_inside_window_and_severity(self):
        tx = [{"date": "2026-09-18", "card_last4": "3079", "description": "GST", "category": "GST", "amount": 60000}]
        b = C.compute_cards(raw(txns=tx), self.TODAY)
        self.assertEqual(len(b["dueItems"]), 1)
        g = b["dueItems"][0]
        self.assertEqual((g["card"], g["date"], g["days"], g["amt"]), ("3079", "2026-10-16", 4, 60000))
        self.assertEqual(g["severity"], "amber")  # 4d > 3d
        self.assertIn("Niloy Debnath", C.due_line(g))

    def test_red_when_3_days_or_overdue(self):
        tx = [{"date": "2026-09-10", "card_last4": "9986", "description": "x", "category": "GST", "amount": 1000}]
        g = C.compute_cards(raw(txns=tx), self.TODAY)["dueItems"][0]
        self.assertLess(g["days"], 0)
        self.assertEqual(g["severity"], "red")
        self.assertIn("overdue", C.due_line(g))

    def test_cleared_spend_never_surfaces(self):
        tx = [{"date": "2026-09-18", "card_last4": "3079", "description": "GST", "category": "GST",
               "amount": 60000, "cleared_on": "2026-10-10"}]
        b = C.compute_cards(raw(txns=tx), self.TODAY)
        self.assertEqual(b["dueItems"], [])
        self.assertIsNone(b["txns"][0]["clearBy"])
        self.assertIsNone(b["cards"][1]["earliestClearBy"])

    def test_outside_window_not_surfaced(self):
        tx = [{"date": "2026-10-11", "card_last4": "9986", "description": "x", "category": "Travel", "amount": 5000}]
        self.assertEqual(C.compute_cards(raw(txns=tx), self.TODAY)["dueItems"], [])

    def test_spends_grouped_per_card_and_clear_by_date(self):
        tx = [{"date": "2026-09-18", "card_last4": "3079", "description": "a", "category": "GST", "amount": 60000},
              {"date": "2026-09-19", "card_last4": "3079", "description": "b", "category": "GST", "amount": 40000}]
        b = C.compute_cards(raw(txns=tx), self.TODAY)
        self.assertEqual(len(b["dueItems"]), 1)
        self.assertEqual(b["dueItems"][0]["amt"], 100000)
        self.assertEqual(b["dueItems"][0]["n"], 2)

    def test_gateway_fee_flags(self):
        b = C.compute_cards(raw(), self.TODAY)
        self.assertFalse(b["statutoryPlan"]["feeVerified"])
        self.assertTrue(any("gateway fee" in w for w in b["warnings"]))
        b2 = C.compute_cards(raw(config={"card_gateway_fee_pct": 1.5, "card_reward_pct_effective": 1.25}), self.TODAY)
        self.assertTrue(b2["statutoryPlan"]["feeVerified"])
        self.assertTrue(b2["statutoryPlan"]["feeExceedsReward"])
        b3 = C.compute_cards(raw(config={"card_gateway_fee_pct": 0.5}), self.TODAY)
        self.assertFalse(b3["statutoryPlan"]["feeExceedsReward"])

    def test_stale_as_of_warns(self):
        r = raw()
        for c in r["cards"]:
            c["as_of_date"] = "2026-10-05"
        b = C.compute_cards(r, self.TODAY)
        self.assertEqual(b["asOfAgeDays"], 7)
        self.assertTrue(any("old" in w for w in b["warnings"]))

    def test_weekend_gap_is_not_stale_but_four_days_is(self):
        r = raw()
        for c in r["cards"]:
            c["as_of_date"] = "2026-10-09"           # Friday
        monday = C.compute_cards(r, date(2026, 10, 12))
        self.assertEqual(monday["asOfAgeDays"], 3)
        self.assertFalse(any("old" in w for w in monday["warnings"]))
        self.assertEqual([i for i in C.health_issues(monday, True) if "stale" in i], [])
        tuesday = C.compute_cards(r, date(2026, 10, 13))
        self.assertTrue(any("old" in w for w in tuesday["warnings"]))


class TestFundingOrderPlan(unittest.TestCase):
    TODAY = date(2026, 10, 12)

    def stat(self, name, amt, due, status="projected"):
        return {"name": name, "amount": amt, "due": due, "status": status, "periodFor": "2026-09-01",
                "daysLeft": (date.fromisoformat(due) - self.TODAY).days}

    def test_cards_exhausted_first_then_od(self):
        b = C.compute_cards(raw(), self.TODAY, statutory=[self.stat("GSTR-3B", 2900000, "2026-10-20")])
        row = b["statutoryPlan"]["rows"][0]
        self.assertEqual(sum(x["amt"] for x in row["routes"]), 390000)   # all available limit used
        self.assertEqual(row["odAmt"], 2900000 - 390000)                 # remainder -> OD
        self.assertEqual(b["statutoryPlan"]["cardRoutable"], 390000)

    def test_limit_not_double_counted_across_obligations(self):
        b = C.compute_cards(raw(), self.TODAY, statutory=[
            self.stat("GSTR-3B", 300000, "2026-10-20"), self.stat("TDS", 300000, "2026-11-07")])
        r0, r1 = b["statutoryPlan"]["rows"]
        self.assertEqual(sum(x["amt"] for x in r0["routes"]), 300000)
        self.assertEqual(sum(x["amt"] for x in r1["routes"]), 90000)    # only 90K of limit left
        self.assertEqual(r1["odAmt"], 210000)

    def test_non_card_payable_goes_straight_to_od(self):
        b = C.compute_cards(raw(), self.TODAY, statutory=[self.stat("PF", 100000, "2026-10-15")])
        row = b["statutoryPlan"]["rows"][0]
        self.assertFalse(row["eligible"])
        self.assertEqual(row["routes"], [])
        self.assertEqual(row["odAmt"], 100000)

    def test_filed_and_zero_items_skipped(self):
        b = C.compute_cards(raw(), self.TODAY, statutory=[
            self.stat("GSTR-3B", 500000, "2026-09-20", status="filed"), self.stat("TDS", 0, "2026-10-20")])
        self.assertEqual(b["statutoryPlan"]["rows"], [])

    def test_overdue_item_clear_by_counts_from_today(self):
        b = C.compute_cards(raw(), self.TODAY, statutory=[self.stat("GSTR-3B", 100000, "2026-10-01")])
        route = b["statutoryPlan"]["rows"][0]["routes"][0]
        self.assertGreaterEqual(date.fromisoformat(route["clearBy"]), self.TODAY)


class TestStatementCycleNotSet(unittest.TestCase):
    """Real baseline: statement day / due gap unknown until Sonal enters them from the HDFC statement."""
    TODAY = date(2026, 10, 12)

    def blank_cycle(self, txns=None):
        r = raw(txns=txns or [])
        for c in r["cards"]:
            c.pop("statement_day"); c.pop("due_gap_days")
        return C.compute_cards(r, self.TODAY)

    def test_no_made_up_statement_dates(self):
        b = self.blank_cycle()
        for c in b["cards"]:
            self.assertIsNone(c["stmtDay"]); self.assertIsNone(c["dueGapDays"])
            self.assertIsNone(c["nextStatement"]); self.assertIsNone(c["nextDue"])
        self.assertTrue(any("statement_day" in w for w in b["warnings"]))

    def test_clear_by_falls_back_to_the_48_day_rule(self):
        d, basis = C.clear_by({"stmtDay": None, "dueGapDays": None}, date(2026, 10, 7), A)
        self.assertEqual(d, date(2026, 11, 24))
        self.assertIn("statement day not set", basis)

    def test_spend_still_gets_a_clear_by(self):
        b = self.blank_cycle([{"date": "2026-09-20", "card_last4": "9986", "description": "x", "category": "GST", "amount": 5000}])
        self.assertEqual(b["txns"][0]["clearBy"], "2026-11-07")
        self.assertEqual(b["dueItems"], [])

    def test_health_reports_missing_cycle(self):
        issues = C.health_issues(self.blank_cycle(), True)
        self.assertTrue(any("statement_day" in i for i in issues))

    def test_clean_baseline_has_no_demo_and_no_due_items(self):
        r = raw()
        b = C.compute_cards(r, self.TODAY)
        self.assertFalse(b["isDemo"])
        self.assertEqual(b["dueItems"], [])


class TestSheetReader(unittest.TestCase):
    def wb(self, cards_rows=None, txn_rows=None, cfg_rows=None, tabs=True):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "README"
        if tabs:
            ws = wb.create_sheet("Credit_Cards")
            ws.append(["card_last4", "card_holder", "product", "credit_limit", "outstanding", "available",
                       "reward_points", "statement_day", "due_gap_days", "activated_on", "welcome_spend_target",
                       "welcome_deadline", "as_of_date", "updated_by", "notes"])
            for r in (cards_rows if cards_rows is not None else [
                    [9986, "Monica Debnath", "HDFC BizBlack Metal", 400000, 138000, None, 1725, 15, 25,
                     datetime(2026, 9, 15), 150000, datetime(2026, 12, 27), datetime(2026, 10, 12), "DEMO", "DEMO DATA"]]):
                ws.append(r)
            wt = wb.create_sheet("Card_Txns")
            wt.append(["date", "card_last4", "description", "category", "amount", "linked_obligation",
                       "challan_ref", "points_earned", "notes", "cleared_on"])
            for r in (txn_rows or []):
                wt.append(r)
            wc = wb.create_sheet("Card_Config")
            wc.append(["key", "value", "note"])
            for r in (cfg_rows or [["card_reward_pct_effective", 1.25, ""]]):
                wc.append(r)
        return wb

    def test_missing_tabs_is_not_an_error(self):
        out = R.parse_cards(self.wb(tabs=False))
        self.assertFalse(out["present"])
        self.assertEqual(out["cards"], [])

    def test_numeric_last4_is_zero_padded_and_dates_iso(self):
        out = R.parse_cards(self.wb(cards_rows=[[247, "X", "p", 1, 0, None, 0, 15, 25, None, 0, None, datetime(2026, 10, 12), "S", ""]]))
        self.assertEqual(out["cards"][0]["card_last4"], "0247")
        self.assertEqual(out["cards"][0]["as_of_date"], "2026-10-12")

    def test_full_card_number_is_rejected_never_stored(self):
        pan = "1111222233334444"   # obviously fake 16-digit value
        out = R.parse_cards(self.wb(cards_rows=[[pan, "X", "p", 1, 0, None, 0, 15, 25, None, 0, None, None, "S", ""]]))
        self.assertEqual(out["cards"], [])
        self.assertTrue(any("4 digits" in w for w in out["meta"]["warnings"]))
        self.assertNotIn(pan, json.dumps(out))
        self.assertFalse(out["present"])

    def test_txn_cleared_on_and_zero_amount_rows(self):
        out = R.parse_cards(self.wb(txn_rows=[
            [datetime(2026, 10, 7), 9986, "TDS", "Income tax", 138000, "TDS Sep-26", None, 1725, "", datetime(2026, 10, 20)],
            [datetime(2026, 10, 8), 9986, "blank", "GST", 0, None, None, 0, "", None]]))
        self.assertEqual(len(out["txns"]), 1)
        self.assertEqual(out["txns"][0]["cleared_on"], "2026-10-20")
        self.assertEqual(out["txns"][0]["card_last4"], "9986")

    def test_registered_in_tab_parsers(self):
        self.assertIn(("Credit_Cards", "sheet_cards.json", R.parse_cards), R.TAB_PARSERS)
        self.assertNotIn("Credit_Cards", R.REQUIRED_TABS)   # cards are optional — never block the sync


class TestTextHelpers(unittest.TestCase):
    TODAY = date(2026, 10, 12)

    def block(self):
        tx = [{"date": "2026-09-18", "card_last4": "3079", "description": "GST", "category": "GST", "amount": 60000}]
        return C.compute_cards(raw(txns=tx), self.TODAY)

    def test_pulse_line(self):
        ln = C.pulse_line(self.block())
        self.assertIn("₹4.1L/₹8.0L (51%)", ln)
        self.assertIn("next clear-by 16 Oct (4d)", ln)
        self.assertEqual(C.pulse_line(None), "")

    def test_health_issues(self):
        self.assertTrue(C.health_issues(None, False))
        issues = C.health_issues(self.block(), True)
        self.assertFalse(any("DEMO" in i for i in issues))
        self.assertEqual(issues, [])        # amber (4d), fresh, not demo -> healthy

    def test_health_flags_demo_stale_and_red_due(self):
        r = raw(txns=[{"date": "2026-09-10", "card_last4": "9986", "description": "x", "category": "GST", "amount": 1000}])
        r["cards"][0]["updated_by"] = "DEMO"
        b = C.compute_cards(r, self.TODAY)
        issues = C.health_issues(b, True)
        self.assertTrue(any("DEMO" in i for i in issues))
        self.assertTrue(any("overdue" in i for i in issues))

    def test_delta_silent_when_unchanged_and_speaks_on_change(self):
        b = self.block()
        snap = C.snapshot(b)
        self.assertEqual(C.delta_lines(b, snap), [])
        b2 = json.loads(json.dumps(b))
        b2["totals"]["outstanding"] += 50000
        lines = C.delta_lines(b2, snap)
        self.assertEqual(len(lines), 1)
        self.assertIn("+₹50K", lines[0])

    def test_delta_reports_new_due_item(self):
        b = self.block()
        self.assertTrue(any("Pay HDFC card" in l for l in C.delta_lines(b, {"outstanding": b["totals"]["outstanding"], "dueKeys": []})))


class TestDeltaSpendsAndCleared(unittest.TestCase):
    TODAY = date(2026, 10, 12)
    TX = {"date": "2026-10-01", "card_last4": "9986", "description": "GST payment", "category": "GST", "amount": 40000}

    def blk(self, txns):
        return C.compute_cards(raw(txns=txns), self.TODAY)

    def test_new_spend_reported_even_if_outstanding_unchanged(self):
        base = self.blk([self.TX])
        snap = C.snapshot(base)
        new = dict(self.TX, date="2026-10-12", description="TDS challan", amount=90000)
        lines = C.delta_lines(self.blk([self.TX, new]), snap)
        self.assertEqual(len(lines), 1)
        self.assertIn("new card spend Monica ·9986 ₹90K · TDS challan", lines[0])

    def test_newly_cleared_bill_reported_once(self):
        snap = C.snapshot(self.blk([self.TX]))
        paid = dict(self.TX, cleared_on="2026-10-12")
        lines = C.delta_lines(self.blk([paid]), snap)
        self.assertEqual(lines, ["✅ card bill cleared: ₹40K across 1 spend(s)"])
        self.assertEqual(C.delta_lines(self.blk([paid]), C.snapshot(self.blk([paid]))), [])

    def test_old_snapshot_without_txn_keys_never_floods(self):
        old = {"outstanding": 410000.0, "dueKeys": []}   # snapshot saved before this feature (no txnKeys)
        self.assertEqual(C.delta_lines(self.blk([self.TX, dict(self.TX, amount=1000)]), old), [])

    def test_more_than_three_new_spends_are_summarised(self):
        snap = C.snapshot(self.blk([]))
        many = [dict(self.TX, amount=1000 * (i + 1), description=f"s{i}") for i in range(5)]
        lines = C.delta_lines(self.blk(many), snap)
        self.assertEqual(len(lines), 4)
        self.assertIn("+2 more new card spends", lines[-1])


class TestTemplateContract(unittest.TestCase):
    T = (ROOT / "dashboards" / "dashboard-template.html").read_text()

    def test_ui_markers_present(self):
        for m in ("CCB", "page-cards", "cc-plan", "Card payment plan", "Funding order (policy)",
                  "priCard", "ccStripHtml", "Clear card payments by"):
            self.assertIn(m, self.T, f"template missing credit-card marker {m!r}")

    def test_ui_is_fully_gated_on_fr_cards(self):
        self.assertIn("(FR.cards && FR.cards.present) ? FR.cards : null", self.T)
        self.assertIn("if (CCB) (function(){", self.T)
        self.assertIn("if (FR.cards && FR.cards.present) {", self.T)   # nav item

    def test_no_demo_data_embedded_in_template(self):
        self.assertNotIn("demo-ribbon", self.T)
        self.assertNotRegex(self.T, r"\b\d{4}\s\d{4}\s\d{4}\s\d{4}\b")   # no card-number-shaped strings

    def test_funding_order_and_clear_by_rule_visible(self):
        self.assertIn("exhaust the card limits first, then OD, treasury last", self.T)
        self.assertIn("never later than HDFC", self.T)


class TestSyncSkillsIncludeCards(unittest.TestCase):
    """All four daily syncs must carry credit-card data (Niloy, 2026-10-12)."""

    def skill(self, name):
        return (ROOT / ".claude" / "scheduled-tasks" / name / "SKILL.md").read_text()

    def test_morning_sync_reports_cards(self):
        t = self.skill("first-rain-monday-sync")
        for m in ("5-CARDS", "cards.py line", "cards.py due", "cards.py health", "cards.py mark-sent",
                  "💳 [`cards.py line` output", "cards first → OD → treasury last", "use date + 48 days"):
            self.assertIn(m, t, f"morning SKILL missing {m!r}")

    def test_midday_and_eod_report_card_deltas(self):
        for n in ("first-rain-midday-refresh", "first-rain-eod-refresh"):
            t = self.skill(n)
            self.assertIn("cards.py delta", t, n)
            self.assertIn("cards.py mark-sent", t, n)
            self.assertIn("Cards · FRBIS checked", t, n)

    def test_heal_check_has_card_health_step(self):
        t = self.skill("first-rain-heal-check")
        for m in ("STEP 4D", "cards.py health", "cards block failed to build", "already-in-today"):
            self.assertIn(m, t, f"heal-check SKILL missing {m!r}")

    def test_skills_never_ask_for_full_card_numbers(self):
        for n in ("first-rain-monday-sync", "first-rain-heal-check", "first-rain-midday-refresh", "first-rain-eod-refresh"):
            self.assertNotRegex(self.skill(n), r"\b\d{4}\s\d{4}\s\d{4}\s\d{4}\b", n)

    def test_alerts_tab_health_line_uses_the_dashboard_score(self):
        # the Python briefing uses a simplified score; the Alerts tab must show the real one
        tpl = (ROOT / "dashboards" / "dashboard-template.html").read_text()
        self.assertIn("'🏥 Health ' + healthScore + '/100 · ' + healthGrade", tpl)


class TestNoSecretsInRepoFiles(unittest.TestCase):
    def test_no_pan_or_cvv_in_card_artifacts(self):
        for rel in ("scripts/projects/cards.py", "scripts/projects/read_cashflow_xlsx.py",
                    "_outputs/reka-card-tabs-demo/Credit_Cards.csv", "_outputs/reka-card-tabs-demo/Card_Txns.csv"):
            p = ROOT / rel
            if not p.exists():
                continue
            txt = p.read_text()
            # any 13-19 digit run (card-number shaped, spaces/dashes allowed) must not appear
            self.assertNotRegex(txt, r"\b(?:\d[ -]?){13,19}\b", f"{rel} contains a card-number-shaped digit run")
            self.assertNotRegex(txt, r"(?i)\bcvv\s*[:=]?\s*\d{3}\b", f"{rel} contains a CVV")


if __name__ == "__main__":
    unittest.main()
